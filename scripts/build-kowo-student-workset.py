from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "kowo-26-27"
TASK_OUT = OUT_DIR / "Aufgabenbeschreibung_BLC26_Kurvenpruefung.docx"
FORM_OUT = OUT_DIR / "Arbeitsvorlage_BLC26_Kurvenpruefung.docx"
HELP_OUT = OUT_DIR / "Schnellhilfe_BLC26_Kurvenpruefung.docx"

BLUE = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F5F6"
GRID = "D9D9D9"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), color)


def margins(cell, top=100, start=115, bottom=100, end=115):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tbl_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), GRID)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def run_style(run, bold=False, italic=False, size=None, color=None):
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Aptos")
    r_pr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def setup_doc(footer_text, compact=False):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.2 if compact else 1.45)
    section.bottom_margin = Cm(1.15 if compact else 1.35)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.0 if compact else 9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.06

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(3)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_style(footer.add_run(footer_text + "     Seite "), size=8, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = footer.add_run()
    field_run._r.extend([begin, instr, end])
    return doc


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def body(doc, text, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        run_style(p.add_run(bold_lead), bold=True)
        run_style(p.add_run(text[len(bold_lead):]), italic=italic)
    else:
        run_style(p.add_run(text), italic=italic)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run_style(p.add_run(item))


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run_style(p.add_run(item))


def table(doc, headers, rows, widths=None, font_size=8.2, blank_rows=0):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    borders(result)
    repeat_header(result.rows[0])
    for idx, label in enumerate(headers):
        cell = result.rows[0].cells[idx]
        shade(cell, BLUE)
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_style(p.add_run(label), bold=True, color="FFFFFF", size=font_size)
        if widths:
            cell.width = Cm(widths[idx])
    all_rows = list(rows) + [["Ihre Antwort" if c == 0 and len(headers) == 1 else "" for c in range(len(headers))] for _ in range(blank_rows)]
    for r_idx, values in enumerate(all_rows):
        row = result.add_row()
        for c_idx, value in enumerate(values):
            cell = row.cells[c_idx]
            margins(cell, top=115 if not value else 95, bottom=115 if not value else 95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2:
                shade(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            run_style(p.add_run(str(value)), size=font_size, italic=(str(value) == "Ihre Antwort"), color="666666" if str(value) == "Ihre Antwort" else None)
            if widths:
                cell.width = Cm(widths[c_idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return result


def response(doc, label="Ihre Antwort", lines=3, help_text=None):
    if help_text:
        body(doc, "Hilfe: " + help_text, bold_lead="Hilfe:", italic=True)
    result = doc.add_table(rows=1, cols=1)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(result)
    repeat_header(result.rows[0])
    cell = result.cell(0, 0)
    shade(cell, PALE_GRAY)
    margins(cell, top=110, bottom=110)
    p = cell.paragraphs[0]
    run_style(p.add_run(label), italic=True, color="666666", size=8.5)
    for _ in range(lines):
        run_style(p.add_run("\n"), size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_title(doc, title, intro=None):
    doc.add_paragraph(title, style="Title")
    if intro:
        body(doc, intro)


def create_task_description():
    doc = setup_doc("KoWo 2026/27 · Aufgabenbeschreibung · Version 0.2")
    page_title(
        doc,
        "BLC26 Kurven prüfen verstehen und vermitteln",
        "Sie untersuchen in einer kleinen Gruppe eine ausgewählte BLC26-Kurve. Sie prüfen die Daten und Quellen, erklären den gesamten Verlauf und übertragen die Ergebnisse vorsichtig auf Praxis und Bauen.",
    )
    heading(doc, "Ihr Ergebnis", 1)
    body(doc, "Sie geben einen kompakten Prüfbericht ab. Die Arbeitsvorlage führt Sie Schritt für Schritt durch die Bearbeitung. Halten Sie die Schnellhilfe beim Arbeiten geöffnet. Das ZUSTAND-Zertifikat auf Seite 1 ist zunächst ein Entwurf; die spätere fachliche Freigabe übernimmt die Projektleitung.")
    bullets(doc, [
        "Seite 1: ZUSTAND-Prüfentwurf und Gesamturteil",
        "Seite 2: Kurve, Datenherkunft und Quellenqualität",
        "Seite 3: drei Kurvensegmente und Methodenwechsel",
        "Seite 4: Bedeutung, Aussagegrenzen und Baubezug",
        "Seite 5: Wechselwirkungen und Verbesserungen",
        "Anhang: Quellen, Nachweise und kurze Dokumentation des KI-Einsatzes",
    ])
    heading(doc, "Die drei Kurvensegmente", 1)
    table(doc, ["Segment", "Was ist gemeint"], [
        ["Rekonstruktion", "Vergangene Werte werden aus indirekten Hinweisen oder historischen Daten wissenschaftlich erschlossen."],
        ["Beobachtung", "Werte werden mit Messverfahren erhoben, ausgewertet und veröffentlicht."],
        ["Projektion", "Modelle berechnen mögliche Zukunftsentwicklungen unter bestimmten Annahmen oder Szenarien."],
    ], widths=[4.1, 12.9], font_size=8.5)
    body(doc, "Die Bezeichnungen können bei einzelnen Kurven abweichen. Entscheidend ist, dass Sie die verwendete Methode und ihre Unsicherheit für jedes Segment prüfen.")
    heading(doc, "Pflichtaufgaben", 1)
    numbered(doc, [
        "Beschreiben Sie, was die Kurve zeigt, welche Einheit verwendet wird und was ein steigender oder fallender Verlauf bedeutet.",
        "Kontrollieren Sie drei Werte anhand der Originalquelle, möglichst einen aus jedem vorhandenen Segment.",
        "Beurteilen Sie Herkunft, Methode, wissenschaftliche Qualität und Aktualität der wichtigsten Quellen.",
        "Vergleichen Sie Rekonstruktion, Beobachtung und Projektion. Erklären Sie die Methodenwechsel und die jeweiligen Unsicherheiten.",
        "Formulieren Sie eine Aussage über den gesamten Verlauf. Trennen Sie Messung, Rekonstruktion, Modellannahme und Interpretation.",
        "Erarbeiten Sie mindestens ein Praxis- oder Baubeispiel. Kennzeichnen Sie, welche zusätzlichen regionalen oder lokalen Daten benötigt werden.",
        "Untersuchen Sie mindestens zwei mögliche Verbindungen zu anderen BLC26-Kurven und belegen Sie die Wirkungsrichtung mit Quellen.",
        "Formulieren Sie konkrete Verbesserungsvorschläge für BLC26 oder das GWL-Panel.",
    ])

    doc.add_page_break()
    heading(doc, "Bewertung und Nachweise", 1)
    table(doc, ["Bewertung", "Bedeutung"], [
        ["Grün", "erfüllt und nachvollziehbar"],
        ["Gelb", "teilweise erfüllt oder nur mit Einschränkung"],
        ["Rot", "nicht erfüllt oder wesentlicher Fehler"],
        ["Grau", "nicht prüfbar oder nicht anwendbar"],
    ], widths=[3.3, 13.7], font_size=8.6)
    body(doc, "Jede gelbe, rote oder graue Bewertung wird kurz begründet. Geben Sie bei fachlichen Aussagen möglichst genau an, wo Sie den Beleg gefunden haben: URL oder DOI sowie Seite, Tabelle, Zeile oder Abschnitt.")
    body(doc, "Wenn eine Aufgabe bei Ihrer Kurve nicht passt oder eine Quelle nicht erreichbar ist, tragen Sie nicht anwendbar oder nicht prüfbar ein und begründen dies kurz. Raten Sie keine Antwort.")
    heading(doc, "Wissenschaftlich arbeiten", 1)
    bullets(doc, [
        "Ein ähnlicher Verlauf zweier Kurven beweist noch keinen ursächlichen Zusammenhang.",
        "Eine Projektion ist keine sichere Vorhersage und kein bereits gemessener Zukunftswert.",
        "Ein globaler Wert darf nicht ohne zusätzliche Prüfung auf einen Ort oder ein Gebäude übertragen werden.",
        "Wenn Sie etwas nicht prüfen können, kennzeichnen Sie dies offen. Eine begründete offene Frage ist besser als eine unbelegte Behauptung.",
    ])
    heading(doc, "Zeitplan", 1)
    table(doc, ["Zeit", "Schwerpunkt", "Zwischenziel"], [
        ["Montagnachmittag", "Kurve öffnen, Begriffe klären, Quellen finden, Aufgaben verteilen", "Sie können erklären, was geprüft werden muss."],
        ["Dienstag", "Werte, Datenaufbereitung, Quellen und Methoden prüfen", "Sie können den Weg von der Quelle zur Kurve erklären."],
        ["Mittwoch", "Verlauf deuten, Baubezug und Wechselwirkungen erarbeiten", "Der Bericht ist inhaltlich weitgehend vollständig."],
        ["Donnerstag", "Kreuzcheck, interne Präsentation, Kurve auf das analoge Chart übertragen", "Die überarbeitete Fassung ist eingereicht."],
        ["Freitag", "Gemeinsamer 30-Minuten-Beitrag", "Jede Gruppe vermittelt ihren wichtigsten Befund."],
    ], widths=[3.0, 7.1, 6.9], font_size=7.9)
    heading(doc, "Wenn Sie früher fertig sind", 1)
    body(doc, "Bearbeiten Sie zuerst alle Pflichtfelder. Danach können Sie zusätzliche Werte kontrollieren, weitere Vergleichsquellen oder Kurvenverbindungen untersuchen, das Baubeispiel vertiefen oder nach Rücksprache eine weitere Kurve übernehmen.")

    doc.add_page_break()
    heading(doc, "KI Werkzeuge", 1)
    body(doc, "Sie dürfen KI-Werkzeuge unterstützend verwenden, beispielsweise zum Erklären von Begriffen, Strukturieren von Notizen oder sprachlichen Überarbeiten. KI-Ausgaben sind keine wissenschaftlichen Quellen.")
    bullets(doc, [
        "Prüfen Sie Zahlen, Literaturangaben und fachliche Aussagen anhand zugänglicher Quellen.",
        "Dokumentieren Sie kurz, welches Werkzeug Sie wofür verwendet und wie Sie das Ergebnis kontrolliert haben.",
        "Alle Gruppenmitglieder müssen die abgegebenen Aussagen selbst erklären können.",
    ])
    heading(doc, "Zusammenarbeit", 1)
    body(doc, "Verteilen Sie die Arbeit innerhalb der Gruppe, besprechen Sie aber das Gesamturteil gemeinsam. Benennen Sie möglichst eine Ansprechperson für organisatorische Rückfragen. Offene Fragen sammeln Sie zunächst im vorgesehenen Feld der Arbeitsvorlage.")
    heading(doc, "Präsentationen", 1)
    body(doc, "Für die interne Präsentation am Donnerstag erklären Sie Kurve, Prüfung, wichtigsten Methodenwechsel, Aussagegrenzen, Baubezug und eine Verbindung zu einer anderen Kurve. Am Freitag stellen Sie im gemeinsamen 30-Minuten-Beitrag nur den wichtigsten Befund Ihrer Gruppe vor.")
    heading(doc, "Freiwillige Namensnennung", 1)
    body(doc, "Die Angabe persönlicher Namen ist freiwillig. Der Prüfbericht kann später öffentlich zugänglich gemacht werden. Aus einer fehlenden Namensnennung entstehen keine Nachteile. Ohne freiwillige Namensnennung wird nur die Gruppenkennung verwendet. Tragen Sie keine Namen anderer Personen ohne deren Zustimmung ein.")
    heading(doc, "Vor der Abgabe", 1)
    bullets(doc, [
        "Sind alle Pflichtfelder bearbeitet oder begründet als nicht prüfbar markiert?",
        "Sind alle Zahlen und Aussagen mit nachvollziehbaren Quellen belegt?",
        "Werden die drei Segmente und ihre Methoden klar unterschieden?",
        "Sind globale Aussage und lokaler Baubezug sauber getrennt?",
        "Kann jedes Gruppenmitglied die zentrale Aussage und die wichtigste Einschränkung erklären?",
    ])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(TASK_OUT)


def create_form():
    doc = setup_doc("KoWo 2026/27 · Arbeitsvorlage · Version 0.2")

    page_title(doc, "ZUSTAND Prüfentwurf", "Füllen Sie diese Seite zuletzt aus. Sie fasst Ihre Prüfung zusammen und bleibt bis zur späteren fachlichen Freigabe ein studentischer Entwurf.")
    table(doc, ["Merkmal", "Eintrag"], [
        ["Kurventitel und Kurvenkennung", ""],
        ["GWL-Beitragslink und BLC26-Kurvenlink", ""],
        ["Prüfdatum und Gruppenkennung", ""],
        ["Namen freiwillig", ""],
        ["Vorlagenversion", "0.2 · 28.09.2026"],
    ], widths=[4.0, 13.1], font_size=8.2)
    heading(doc, "Gesamturteil", 1)
    body(doc, "Markieren Sie genau eine Aussage:  ☐ geprüft und plausibel   ☐ geprüft mit Einschränkungen   ☐ nicht ausreichend belegbar   ☐ Korrektur erforderlich")
    table(doc, ["Prüfbereich", "Ampel", "Kurzbegründung"], [
        ["Datenherkunft", "", ""],
        ["Übernahme und Aufbereitung", "", ""],
        ["Wissenschaftliche Qualität und Aktualität", "", ""],
        ["Drei Segmente und Methodenwechsel", "", ""],
        ["Interpretation und Aussagegrenzen", "", ""],
    ], widths=[5.0, 2.1, 10.0], font_size=7.9)
    heading(doc, "Zentrale Aussage", 2)
    response(doc, "Der gesamte Kurvenverlauf bedeutet …", 1, "Fassen Sie Rekonstruktion, Beobachtung und Projektion in höchstens vier Sätzen zusammen.")
    heading(doc, "Wichtigste Einschränkung", 2)
    response(doc, "Beim Lesen der Kurve muss besonders beachtet werden, dass …", 0)
    body(doc, "Status: Studentischer Prüfentwurf. Fachliche Freigabe und Veröffentlichung erfolgen später durch die Projektleitung.", bold_lead="Status:")

    doc.add_page_break()
    heading(doc, "Kurve und Datenherkunft", 1)
    heading(doc, "K 01 Was zeigt die Kurve", 2)
    response(doc, lines=1, help_text="Nennen Sie Größe, Einheit, Raumbezug, Gesamtzeitraum und die Bedeutung eines steigenden oder fallenden Verlaufs. Schreiben Sie höchstens fünf Sätze.")
    heading(doc, "K 02 Grunddaten", 2)
    table(doc, ["Merkmal", "Ihre Angabe"], [
        ["dargestellte Größe und Einheit", ""], ["räumlicher Bezug und Gesamtzeitraum", ""],
        ["Bezugsgröße oder Referenzperiode", ""],
    ], widths=[6.1, 11.0], font_size=8.2)
    heading(doc, "K 03 Kontrollierte Werte", 2)
    table(doc, ["Segment", "Jahr oder Zeitraum", "Wert BLC26", "Wert Quelle", "Abgleich", "Fundstelle"], [
        ["", "", "", "", "", ""], ["", "", "", "", "", ""], ["", "", "", "", "", ""],
    ], widths=[2.5, 2.8, 2.5, 2.5, 2.3, 4.1], font_size=7.3)
    heading(doc, "K 04 Aufbereitung", 2)
    response(doc, lines=1, help_text="Wurden Werte ausgewählt, umgerechnet, geglättet, verbunden oder neu auf eine Referenzperiode bezogen? Ist dies nachvollziehbar?")
    heading(doc, "K 05 Quellenqualität", 2)
    table(doc, ["Prüffrage", "Ampel", "Begründung und Nachweis"], [
        ["Originalquelle und Herausgeber eindeutig", "", ""],
        ["Methode und Aufbereitung nachvollziehbar", "", ""],
        ["Unsicherheiten und Grenzen angegeben", "", ""],
        ["Wissenschaftlich eingeordnet und aktuell genug", "", ""],
    ], widths=[6.3, 2.0, 8.8], font_size=7.7)

    doc.add_page_break()
    heading(doc, "Drei Segmente und Methodenwechsel", 1)
    body(doc, "Hilfe: Rekonstruktion erschließt Vergangenes indirekt, Beobachtung beruht auf Messungen, Projektion berechnet mögliche Zukunftsentwicklungen unter Annahmen.", bold_lead="Hilfe:", italic=True)
    heading(doc, "S 01 Segmentvergleich", 2)
    table(doc, ["Segment", "Zeitraum", "Datentyp und Methode", "Annahmen", "Wichtigste Unsicherheit"], [
        ["1", "", "", "", ""], ["2", "", "", "", ""], ["3", "", "", "", ""],
    ], widths=[1.5, 2.4, 5.3, 3.7, 4.3], font_size=7.4)
    heading(doc, "S 02 Übergang zwischen Segment 1 und 2", 2)
    response(doc, lines=3, help_text="Was ändert sich bei Datenquelle, Methode, Genauigkeit oder zeitlicher Auflösung? Kann ein sichtbarer Sprung durch den Methodenwechsel beeinflusst sein?")
    heading(doc, "S 03 Übergang zwischen Segment 2 und 3", 2)
    response(doc, lines=3, help_text="Wo endet die Messung und wo beginnt die Modellprojektion? Bezeichnet ein dargestelltes Jahr wirklich ein Einzeljahr oder einen Zeitraum?")
    heading(doc, "S 04 Merksatz", 2)
    response(doc, "Ab dem Jahr oder Zeitraum … ändert sich die Methode von … zu … . Das ist wichtig, weil …", 2)
    heading(doc, "S 05 Verständlichkeit der Darstellung", 2)
    table(doc, ["Prüffrage", "Ampel", "Kurzbegründung"], [
        ["Sind die drei Datentypen erkennbar", "", ""],
        ["Sind Übergänge und Überlappungen verständlich", "", ""],
        ["Werden Annahmen und Unsicherheiten ausreichend erklärt", "", ""],
    ], widths=[7.3, 2.0, 7.8], font_size=7.8)

    doc.add_page_break()
    heading(doc, "Bedeutung und Baubezug", 1)
    heading(doc, "B 01 Bedeutung des gesamten Verlaufs", 2)
    response(doc, lines=4, help_text="Beziehen Sie alle drei Segmente ein. Was ist langfristig erkennbar, wo gibt es Brüche oder Wendepunkte und wovon hängt die Zukunft ab? Schreiben Sie höchstens acht Sätze.")
    heading(doc, "B 02 Sicherheit der Aussagen", 2)
    table(doc, ["Aussageart", "Ihre Aussage", "Beleg"], [
        ["gut belegt", "", ""], ["mit Einschränkung", "", ""], ["nicht ableitbar", "", ""],
    ], widths=[3.5, 9.2, 4.4], font_size=7.8)
    heading(doc, "B 03 Praxis oder Baubeispiel", 2)
    response(doc, lines=4, help_text="Beschreiben Sie ein konkretes Beispiel, eine betroffene Entscheidung und mindestens eine zusätzliche Fachquelle. Kennzeichnen Sie, ob der Zusammenhang belegt oder von Ihnen begründet übertragen wurde.")
    heading(doc, "B 04 Zusätzliche Informationen", 2)
    response(doc, lines=2, help_text="Welche regionalen, lokalen oder gebäudespezifischen Daten wären für eine konkrete Entscheidung erforderlich?")
    heading(doc, "B 05 Gefahr einer Fehlinterpretation", 2)
    response(doc, "Wenn die Kurve zu einfach interpretiert wird, könnte …", 2)

    doc.add_page_break()
    heading(doc, "Wechselwirkungen und Verbesserungen", 1)
    heading(doc, "W 01 Verbindungen zu anderen Kurven", 2)
    body(doc, "Hilfe: Unterscheiden Sie Ursache, Wirkung, Rückkopplung, gemeinsame Ursache und bloße zeitliche Übereinstimmung. Ähnliche Verläufe allein beweisen keine Ursache.", bold_lead="Hilfe:", italic=True)
    table(doc, ["Andere Kurve", "Verbindung", "Art und Richtung", "Beleg", "Sicherheit hoch mittel niedrig"], [
        ["", "", "", "", ""], ["", "", "", "", ""],
    ], widths=[3.0, 5.7, 3.4, 3.2, 2.2], font_size=7.3)
    heading(doc, "W 02 Wirkungsskizze", 2)
    response(doc, "Skizzieren Sie Ursachen, Wirkungen und mögliche Rückkopplungen mit Pfeilen.", 7)
    heading(doc, "V 01 Verbesserungsvorschläge", 2)
    table(doc, ["Ziel", "Beobachtung", "Vorgeschlagene Verbesserung", "Dringlichkeit", "Beleg"], [
        ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""],
    ], widths=[2.1, 4.6, 6.3, 2.2, 2.5], font_size=7.1)
    heading(doc, "V 02 Wichtigste Empfehlung", 2)
    response(doc, "Die wichtigste Verbesserung wäre …", 2)

    doc.add_page_break()
    heading(doc, "Quellen Nachweise und Arbeitsdokumentation", 1)
    heading(doc, "Q 01 Quellenverzeichnis", 2)
    body(doc, "Nennen Sie möglichst Originalquellen. Geben Sie Autor oder Institution, Titel, Jahr, DOI oder URL und das Abrufdatum an.")
    table(doc, ["Nr", "Vollständige Quellenangabe", "Verwendete Stelle"], [
        ["1", "", ""], ["2", "", ""], ["3", "", ""], ["4", "", ""],
    ], widths=[1.3, 11.1, 4.7], font_size=7.7)
    heading(doc, "D 01 KI Einsatz", 2)
    table(doc, ["Werkzeug", "Wofür verwendet", "Wie überprüft"], [["", "", ""]], widths=[3.2, 6.8, 7.1], font_size=7.8)
    heading(doc, "D 02 Offene Fragen an die Lehrperson", 2)
    response(doc, lines=2)
    heading(doc, "D 03 Kreuzcheck durch eine andere Gruppe", 2)
    table(doc, ["Prüfung", "Ja", "Noch offen", "Hinweis"], [
        ["Zentrale Aussage verständlich", "☐", "☐", ""],
        ["Quellen und Fundstellen nachvollziehbar", "☐", "☐", ""],
        ["Drei Segmente sauber unterschieden", "☐", "☐", ""],
        ["Baubezug als Übertragung gekennzeichnet", "☐", "☐", ""],
        ["Kurvenverbindungen nicht nur zeitlich begründet", "☐", "☐", ""],
    ], widths=[8.0, 1.6, 2.4, 5.1], font_size=7.7)
    heading(doc, "D 04 Freiwillige Namensnennung", 2)
    body(doc, "Die Nennung persönlicher Namen ist freiwillig. Ohne Namensnennung wird nur die Gruppenkennung verwendet. Eine spätere Veröffentlichung eingetragener Namen erfolgt erst nach gesonderter Bestätigung.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(FORM_OUT)


def create_quick_help():
    doc = setup_doc("KoWo 2026/27 · Schnellhilfe · Version 0.2", compact=True)
    page_title(doc, "Schnellhilfe zur BLC26 Kurvenprüfung", "Nutzen Sie diese Seite, wenn Sie nicht wissen, wie Sie beginnen oder was ein Begriff bedeutet. Offen gebliebene Fragen sind erlaubt, wenn Sie diese nachvollziehbar dokumentieren.")
    heading(doc, "So beginnen Sie", 1)
    numbered(doc, [
        "Öffnen Sie den GWL-Beitrag, die BLC26-Kurve und die dort genannte Originalquelle.",
        "Schreiben Sie in einem Satz auf, was auf der x-Achse und der y-Achse dargestellt wird.",
        "Markieren Sie, welche Teile Rekonstruktion, Beobachtung und Projektion sind. Fehlt ein Teil, notieren Sie dies.",
        "Prüfen Sie danach drei Werte und bearbeiten Sie die Arbeitsvorlage der Reihe nach. Seite 1 füllen Sie zuletzt aus.",
    ])
    heading(doc, "Wichtige Begriffe", 1)
    table(doc, ["Begriff", "Einfache Bedeutung"], [
        ["Originalquelle", "Die Stelle, an der Daten oder Studienergebnisse zuerst veröffentlicht wurden."],
        ["Fundstelle", "Die genaue Stelle im Beleg, zum Beispiel Seite 12, Tabelle 3 oder Datenzeile 2020."],
        ["Referenzperiode", "Ein Vergleichszeitraum, auf den Werte bezogen werden, zum Beispiel der Mittelwert 1850 bis 1900."],
        ["Aufbereitung", "Was mit Rohdaten gemacht wurde, etwa auswählen, mitteln, umrechnen, glätten oder verbinden."],
        ["Unsicherheit", "Der bekannte Bereich, in dem ein Wert oder Ergebnis wahrscheinlich liegt."],
        ["Projektion", "Eine berechnete mögliche Zukunft unter bestimmten Annahmen, keine sichere Vorhersage."],
        ["wissenschaftlich eingeordnet", "Methode, Grenzen und Bedeutung wurden fachlich beschrieben oder geprüft."],
    ], widths=[4.5, 12.6], font_size=7.8)
    heading(doc, "Wenn etwas nicht funktioniert", 1)
    table(doc, ["Situation", "So gehen Sie vor"], [
        ["Quelle nicht erreichbar", "URL und Zugriffsversuch notieren, grau bewerten und die Lehrperson fragen."],
        ["Werte stimmen nicht überein", "Beide Werte und Fundstellen festhalten; Einheit, Zeitraum und Rundung prüfen; nichts stillschweigend korrigieren."],
        ["Kurve hat nicht drei Segmente", "Vorhandene Segmente bearbeiten und fehlende als nicht anwendbar kennzeichnen."],
        ["Baubezug ist nicht direkt belegt", "Als begründete Übertragung kennzeichnen und zusätzliche Fachquelle nennen."],
        ["Sie sind unsicher", "Unsicherheit benennen, bisherigen Prüfweg dokumentieren und eine konkrete Frage formulieren."],
    ], widths=[4.8, 12.3], font_size=7.8)
    heading(doc, "Vor jeder Bewertung", 1)
    body(doc, "Trennen Sie immer drei Dinge: Was steht in der Quelle? Was zeigt die BLC26-Kurve? Was schließen Sie selbst daraus? Eine ähnliche zeitliche Entwicklung zweier Kurven beweist noch keine Ursache.")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(HELP_OUT)


if __name__ == "__main__":
    create_task_description()
    create_form()
    create_quick_help()
    print(TASK_OUT)
    print(FORM_OUT)
    print(HELP_OUT)
