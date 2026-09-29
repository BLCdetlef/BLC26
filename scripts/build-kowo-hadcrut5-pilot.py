from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "kowo-26-27" / "Beispiel_Pruefbericht_HadCRUT5.docx"

BLUE = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F5F6"
GRID = "D9D9D9"
GREEN = "2E7D32"
YELLOW = "A66A00"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), GRID)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def style_run(run, bold=False, color=None, size=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        style_run(p.add_run(bold_lead), bold=True)
        style_run(p.add_run(text[len(bold_lead):]))
    else:
        style_run(p.add_run(text))
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_run(p.add_run(item))


def add_table(doc, headers, rows, widths=None, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    repeat_header(hdr)
    for idx, label in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_fill(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(label), bold=True, color="FFFFFF", size=font_size)
        if widths:
            cell.width = Cm(widths[idx])
    for r_idx, values in enumerate(rows):
        row = table.add_row()
        for c_idx, value in enumerate(values):
            cell = row.cells[c_idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2:
                set_cell_fill(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(str(value)), size=font_size)
            if widths:
                cell.width = Cm(widths[c_idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(paragraph.add_run("Seite "), size=8, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([fld_char1, instr_text, fld_char2])


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.45)
section.bottom_margin = Cm(1.35)
section.left_margin = Cm(1.65)
section.right_margin = Cm(1.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.2)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_p_pr = title_style.element.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)

for name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(7)
    style.paragraph_format.space_after = Pt(3)

footer = section.footer
footer_p = footer.paragraphs[0]
style_run(footer_p.add_run("KoWo 2026/27 · Beispielbericht HadCRUT5 · Kein Zertifikat     "), size=8, color="666666")
add_page_number(footer_p)

# Seite 1
p = doc.add_paragraph("ZUSTAND Beispielprüfbericht", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
sub = doc.add_paragraph()
style_run(sub.add_run("Globale Temperaturänderung 1700 bis 2100"), bold=True, size=14)
add_body(doc, "Beispiel für die Kompaktwoche 2026/27. Dieser vollständig ausgefüllte Bericht zeigt, wie Ergebnisse knapp, nachvollziehbar und mit Einschränkungen begründet werden können. Er ist kein Zertifikat und darf nicht unverändert als eigene Prüfleistung übernommen werden.")

add_table(doc, ["Merkmal", "Eintrag"], [
    ["Kurve", "PAGES2k Rekonstruktion · HadCRUT5 Beobachtung · IPCC AR6 Projektionen"],
    ["Größe", "Änderung der globalen bodennahen Oberflächentemperatur gegenüber 1850–1900"],
    ["Zeitraum", "1700–1849 · 1850–2025 · 2021–2100"],
    ["Gruppe", "Beispielgruppe · keine reale studentische Einreichung"],
    ["Namen", "nicht angegeben · Namensnennung ist freiwillig"],
    ["Vorlage", "Beispielversion 1.0 · 29.09.2026"],
], widths=[4.0, 13.2], font_size=8.6)

add_heading(doc, "Gesamturteil des Beispiels", 1)
add_body(doc, "Fachlich plausibel mit dokumentierten Einschränkungen. Herkunft, Auswahl und Umrechnung der drei Segmente sind nachvollziehbar. HadCRUT5 5.2 hat die im BLC26 verwendete Version 5.1 am 18.09.2026 abgelöst; deshalb bleibt die Aktualität gelb und eine spätere Datenaktualisierung ist zu prüfen.")

add_table(doc, ["Prüfbereich", "Ampel", "Kurzbegründung"], [
    ["Datenherkunft", "Grün", "Primärdaten und Fundstellen sind benannt."],
    ["Übernahme und Aufbereitung", "Grün", "Referenzperioden und Auswahlregeln sind dokumentiert; Stichprobe stimmt."],
    ["Wissenschaftliche Qualität", "Grün", "NOAA, Met Office und IPCC; Methoden und Unsicherheiten publiziert."],
    ["Methodenwechsel", "Gelb", "Fachlich nachvollziehbar, in der Nutzung aber erklärungsbedürftig."],
    ["Aktualität", "Gelb", "BLC26 nutzt HadCRUT5 5.1; Met Office führt inzwischen 5.2."],
    ["Gesamtnachvollziehbarkeit", "Grün", "Aussage ist mit den genannten Einschränkungen plausibel."],
], widths=[4.2, 2.0, 11.0], font_size=8.0)

add_heading(doc, "Zentrale Aussage", 2)
add_body(doc, "Die Rekonstruktion zeigt vor 1850 Schwankungen nahe der Referenzperiode. Die Messreihe zeigt besonders seit der zweiten Hälfte des 20. Jahrhunderts eine starke Erwärmung. Die Projektionen zeigen, dass die weitere Entwicklung wesentlich vom künftigen Emissionspfad abhängt.")
add_body(doc, "Status: Veröffentlichter Beispielbericht, kein Zertifikat. Die Ampeln und Formulierungen dienen der Orientierung; studentische Gruppen müssen ihre Prüfung eigenständig durchführen und belegen.", bold_lead="Status:")

# Seite 2
doc.add_page_break()
add_heading(doc, "Kurve und Datenherkunft", 1)
add_heading(doc, "K 01 Was zeigt die Kurve", 2)
add_body(doc, "Die Kurve zeigt die Änderung der globalen bodennahen Oberflächentemperatur in Grad Celsius gegenüber dem Mittel 1850–1900. Positive Werte bedeuten eine wärmere globale Mitteltemperatur als in dieser Referenzperiode. Die Darstellung verbindet Rekonstruktion, Beobachtung und drei Zukunftsszenarien.")

add_heading(doc, "K 02 Stichprobe der Beobachtungswerte", 2)
add_table(doc, ["Jahr", "Original 1961–1990", "Mittel 1850–1900", "BLC26 neu bezogen", "Ergebnis"], [
    ["1850", "−0,4265 °C", "−0,357361 °C", "−0,0691 °C", "stimmt"],
    ["1900", "−0,2302 °C", "−0,357361 °C", "+0,1272 °C", "stimmt"],
    ["1950", "−0,2397 °C", "−0,357361 °C", "+0,1177 °C", "stimmt"],
    ["2000", "+0,3255 °C", "−0,357361 °C", "+0,6829 °C", "stimmt"],
    ["2025", "+1,0480 °C", "−0,357361 °C", "+1,4054 °C", "stimmt"],
], widths=[1.7, 3.6, 3.7, 3.8, 2.2], font_size=7.8)
add_body(doc, "Kontrollrechnung: Vom veröffentlichten HadCRUT5-Wert wird das arithmetische Mittel 1850–1900 der verwendeten Reihe abgezogen. Die Rechnung wurde anhand der offiziellen CSV unabhängig wiederholt. Die 95-%-Grenzen werden um denselben konstanten Wert verschoben.")

add_heading(doc, "K 03 Aufbereitung im BLC26", 2)
add_bullets(doc, [
    "Alle 176 vollständigen Jahreswerte 1850–2025 sind im Import enthalten; das noch unvollständige Jahr 2026 bleibt ausgeschlossen.",
    "Für die Anzeige wird ausgedünnt, aber nicht interpoliert: erster und letzter Punkt bleiben erhalten, dazwischen liegen mindestens fünf Jahre.",
    "PAGES2k und HadCRUT5 werden jeweils mit ihrem eigenen Mittel 1850–1900 auf dieselbe Bezugsperiode gebracht.",
    "IPCC-Werte werden nicht umgerechnet; die dargestellten Jahre 2030, 2050 und 2090 stehen für 20-Jahres-Mittel.",
])

add_heading(doc, "K 04 Quellenqualität", 2)
add_table(doc, ["Quelle", "Einordnung", "Prüfurteil"], [
    ["PAGES2k 2019 · NOAA/NCEI", "Offener Forschungsdatensatz mit DOI und Ensemblebereich", "Grün"],
    ["HadCRUT5 · Met Office und UEA", "Begutachtete Methodik; 200 Ensemble-Realisierungen bilden Unsicherheiten ab", "Grün"],
    ["IPCC AR6 WGI", "Bewertete Synthese mehrerer Modelle und Evidenzlinien", "Grün"],
    ["Versionsstand HadCRUT5", "5.1 ist reproduzierbar; 5.2 ist inzwischen aktuell", "Gelb"],
], widths=[5.0, 9.7, 2.2], font_size=8.0)

# Seite 3
doc.add_page_break()
add_heading(doc, "Drei Segmente und Methodenwechsel", 1)
add_heading(doc, "S 01 Segmentvergleich", 2)
add_table(doc, ["Segment", "Zeitraum", "Methode", "Wichtigste Unsicherheit"], [
    ["Rekonstruktion", "1700–1849", "Median eines methodenübergreifenden PAGES2k-Ensembles mit 7000 Mitgliedern", "Proxybasierte Rekonstruktion; breiter 95-%-Bereich, keine direkte Thermometermessung"],
    ["Beobachtung", "1850–2025", "HadCRUT5 kombiniert Landluft- und Meeresoberflächentemperaturen und statistische räumliche Ergänzung", "Mess-, Stichproben-, Homogenisierungs- und Abdeckungsunsicherheit"],
    ["Projektion", "2021–2100", "IPCC-Bestschätzungen für SSP1-1.9, SSP2-4.5 und SSP5-8.5; 20-Jahres-Mittel", "Abhängig von Szenario, Modellantwort und natürlicher Variabilität; keine Einzeljahresvorhersage"],
], widths=[3.0, 2.6, 6.6, 5.2], font_size=7.9)

add_heading(doc, "S 02 Übergang von Rekonstruktion zu Beobachtung", 2)
add_body(doc, "Zwischen 1849 und 1850 wechselt die Datenbasis von indirekten Klima-Proxys und statistischen Rekonstruktionsmethoden zu instrumentellen Land- und Meeresbeobachtungen. Beide Segmente sind auf 1850–1900 bezogen, wurden dafür aber getrennt neu zentriert. Ein Unterschied am Übergang darf daher nicht als exakter Temperatursprung gelesen werden.")

add_heading(doc, "S 03 Übergang von Beobachtung zu Projektion", 2)
add_body(doc, "Ab 2021 überlappen Messzeitraum und Projektionszeiträume. Die Projektionspunkte bezeichnen Mittelwerte von 20 Jahren und keine Messung im Mittelpunktjahr. Die nahe Zukunft unterscheidet sich zwischen den Szenarien zunächst wenig; bis 2081–2100 weichen die Bestschätzungen deutlich auseinander.")

add_heading(doc, "S 04 Merksatz", 2)
add_body(doc, "Ab 1850 ändert sich die Methode von Rekonstruktion zu instrumenteller Beobachtung. Für die Zukunft ändert sie sich zu szenariobasierter Modellprojektion. Deshalb bilden die drei Segmente eine gemeinsame Temperaturgröße ab, besitzen aber unterschiedliche Evidenz, zeitliche Auflösung und Unsicherheit.")

add_heading(doc, "S 05 Verständlichkeit der Darstellung", 2)
add_table(doc, ["Prüffrage", "Ampel", "Kurzbegründung"], [
    ["Sind die drei Datentypen unterscheidbar", "Grün", "Segmente und Szenarien sind in den technischen Details bezeichnet."],
    ["Ist die zeitliche Überlappung verständlich", "Gelb", "2021–2025 sind Messwerte vorhanden, während der erste Projektionspunkt ein Mittel 2021–2040 bezeichnet."],
    ["Ist der Methodenwechsel sofort sichtbar", "Gelb", "Die Erklärung ist vorhanden, könnte aber direkt an den Übergängen deutlicher erscheinen."],
], widths=[7.0, 2.0, 7.6], font_size=8.1)

# Seite 4
doc.add_page_break()
add_heading(doc, "Bedeutung und Baubezug", 1)
add_heading(doc, "B 01 Bedeutung des gesamten Verlaufs", 2)
add_body(doc, "Die Rekonstruktion ordnet die vorindustriellen Schwankungen ein. Die instrumentelle Reihe zeigt eine deutliche Erwärmung, die in den jüngsten Jahren weit über den Schwankungsbereich des historischen Segments hinausreicht. Die Zukunftssegmente zeigen keinen feststehenden Verlauf, sondern unterschiedliche mögliche Erwärmungsniveaus unter verschiedenen Emissionspfaden.")

add_heading(doc, "B 02 Sicherheit der Aussagen", 2)
add_table(doc, ["Aussageart", "Beispiel"], [
    ["Gut belegt", "Die globale Mitteltemperatur ist gegenüber 1850–1900 deutlich gestiegen; HadCRUT5 weist 2025 rund +1,41 °C aus."],
    ["Mit Einschränkung", "Die drei ausgewählten SSPs zeigen eine Spannweite möglicher langfristiger Entwicklungen; sie sind keine Wahrscheinlichkeitsrangfolge."],
    ["Nicht ableitbar", "Die Kurve sagt nicht, welche Temperatur an einem bestimmten Ort oder in einem einzelnen Gebäude auftreten wird."],
], widths=[4.0, 13.0], font_size=8.3)

add_heading(doc, "B 03 Beispiel für den Baubezug", 2)
add_body(doc, "Steigende globale Temperaturen erhöhen in vielen Regionen den Kühlbedarf und das Risiko sommerlicher Überhitzung. Für Gebäude können daraus Anforderungen an Verschattung, Fensterflächen, Speichermasse, Nachtlüftung, Begrünung und effiziente Kühlung entstehen. Zugleich kann in kalten Regionen der Heizbedarf sinken. Welche Maßnahme an einem konkreten Standort sinnvoll ist, muss mit regionalen Klimaprojektionen und gebäudespezifischen Randbedingungen geprüft werden.")

add_heading(doc, "B 04 Konkrete Arbeitsfrage", 2)
add_body(doc, "Beispiel: Reicht der bisherige sommerliche Wärmeschutz eines geplanten Gebäudes auch für künftige Lübecker Sommer? Die globale Kurve liefert den Hintergrund, aber noch keine Bemessungsdaten. Benötigt werden mindestens regionale Klimadaten, Nutzung, Bauweise, Verschattung, interne Lasten und eine geeignete Überhitzungsberechnung.")

add_heading(doc, "B 05 Gefahr einer Fehlinterpretation", 2)
add_body(doc, "Aus +1,41 °C globaler Temperaturänderung darf nicht geschlossen werden, dass jeder Ort oder jeder Sommertag genau 1,41 °C wärmer ist. Globale Jahresmittel, regionale Veränderungen und Extremereignisse sind unterschiedliche Größen.")

add_heading(doc, "Auftrag an die Studierenden", 2)
add_body(doc, "Im echten Gruppenbericht wird hier ein selbst recherchiertes Bau- oder Praxisbeispiel mit mindestens einer zusätzlichen Fachquelle ergänzt. Der Zusammenhang muss als belegt, begründet übertragen oder noch offen gekennzeichnet werden.")

# Seite 5
doc.add_page_break()
add_heading(doc, "Wechselwirkungen und Verbesserungen", 1)
add_heading(doc, "W 01 Verbindungen zu anderen Kurven", 2)
add_table(doc, ["Andere Kurve", "Mögliche Verbindung", "Art", "Sicherheit"], [
    ["Atmosphärisches CO₂", "Steigende Treibhausgaskonzentrationen erhöhen den Strahlungsantrieb und tragen zur Erwärmung bei.", "Ursache und Wirkung", "gut belegt"],
    ["Globaler Meeresspiegel", "Erwärmung führt unter anderem zu thermischer Ausdehnung und verstärkt Eisschmelze; beides erhöht den Meeresspiegel.", "Wirkungspfad", "gut belegt"],
    ["Arktische Meereisfläche", "Erwärmung begünstigt Meereisverlust; geringere Rückstrahlung kann die Erwärmung regional verstärken.", "Wirkung und Rückkopplung", "gut belegt, regional unterschiedlich"],
], widths=[3.5, 8.8, 3.0, 2.4], font_size=7.8)
add_body(doc, "Prüfregel: Zeitliche Ähnlichkeit allein reicht nicht. Jede Wirkungsrichtung braucht eine Quelle; gemeinsame Ursachen und Rückkopplungen sind getrennt zu benennen.")

add_heading(doc, "V 01 Verbesserungsvorschläge", 2)
add_table(doc, ["Ziel", "Beobachtung", "Vorschlag", "Dringlichkeit"], [
    ["BLC26", "HadCRUT5 5.2 ist seit 18.09.2026 verfügbar.", "Version 5.2 gegen 5.1 prüfen; Änderung dokumentieren oder 5.1 bewusst festschreiben.", "hoch"],
    ["BLC26", "Der Projektionspunkt 2030 ist ein Mittel 2021–2040.", "Periodenbezug auch direkt am Punkt und in barrierefreier Beschriftung anzeigen.", "hoch"],
    ["BLC26", "Methodenwechsel sind in Details beschrieben, aber nicht als eigene Übergänge geführt.", "Übergänge 1849/1850 und Beobachtung/Projektion sichtbar markieren.", "mittel"],
    ["GWL", "Globale Kurve kann als lokale Vorhersage missverstanden werden.", "Hinweis auf erforderliche regionale Klimadaten im Baubezug ergänzen.", "mittel"],
], widths=[2.2, 5.4, 7.5, 2.2], font_size=7.6)

add_heading(doc, "V 02 Wichtigste Empfehlung", 2)
add_body(doc, "Die drei Datentypen sollten nicht nur farblich oder in technischen Details unterschieden werden. Die Darstellung sollte direkt erklären, dass Rekonstruktion, Jahresmessung und 20-Jahres-Szenariomittel methodisch verschieden sind.")

add_heading(doc, "D 01 Dokumentation des KI Einsatzes", 2)
add_table(doc, ["Verwendung", "Notwendige Kontrolle"], [
    ["Strukturierung des Beispielberichts und Formulierung kurzer Musterantworten", "Zahlen gegen BLC26 und Originaldatei geprüft; wissenschaftliche Aussagen über Primärquellen eingeordnet"],
], widths=[8.0, 9.0], font_size=8.2)

# Seite 6
doc.add_page_break()
add_heading(doc, "Quellen und Hinweise zum Beispielbericht", 1)
add_heading(doc, "Verwendete Hauptquellen", 2)
sources = [
    ("NOAA NCEI PAGES2k 2019 Global Common Era Temperature Reconstructions", "https://doi.org/10.25921/tkxp-vn12"),
    ("Met Office HadCRUT5 Datensatz und Methodik", "https://hadleyserver.metoffice.gov.uk/hadcrut5/"),
    ("HadCRUT5 5.1 globale Jahresreihe", "https://www.metoffice.gov.uk/hadobs/hadcrut5/data/HadCRUT.5.1.0.0/analysis/diagnostics/HadCRUT.5.1.0.0.analysis.summary_series.global.annual.csv"),
    ("Met Office Versionshinweise zu HadCRUT5 5.2", "https://hadleyserver.metoffice.gov.uk/hadcrut5/data/versions/HadCRUT.5.2.0.0_release_notes.html"),
    ("IPCC AR6 WGI Kapitel 4 mit bewerteten SSP Temperaturänderungen", "https://www.ipcc.ch/report/ar6/wg1/chapter/chapter-4/"),
    ("IPCC AR6 WGIII Kapitel 9 Gebäude", "https://www.ipcc.ch/report/ar6/wg3/chapter/chapter-9/"),
    ("BLC26 Kurve Globale Temperaturänderung", "https://blcdetlef.github.io/BLC26/?curve=knowledge%3Adata%2Fknowledge%2Fgwl_climate_temperature_global_v0.2.json%23global_temperature_hadcrut5_1850_2025"),
    ("GWL Beitrag Globale Temperaturentwicklung", "https://blcdetlef.github.io/gwl-panel/?boundary=climate&item=global-temperature"),
]
for label, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    add_hyperlink(p, label, url)

add_heading(doc, "Was dieses Beispiel zeigt", 2)
add_bullets(doc, [
    "Kurze Antworten reichen aus, wenn jede Bewertung durch eine konkrete Fundstelle oder Kontrollrechnung gestützt wird.",
    "Die Segmenttabelle ist der gemeinsame Kern, weil sie Rekonstruktion, Messung und Projektion sowie deren Methodenwechsel sichtbar trennt.",
    "Aktualität ist ein eigener Prüfpunkt: Ein reproduzierbarer Datensatz kann fachlich korrekt und dennoch nicht mehr die jüngste Version sein.",
    "Der Baubezug ist eine begründete Übertragung; globale Temperaturwerte sind keine lokalen Bemessungsdaten für ein Gebäude.",
    "Bei Wechselwirkungen sind Ursache, Wirkung, Rückkopplung, gemeinsame Ursache und bloße zeitliche Ähnlichkeit voneinander zu unterscheiden.",
])

add_heading(doc, "Hinweis für die eigene Bearbeitung", 2)
add_body(doc, "Nutzen Sie Aufbau und Begründungstiefe dieses Beispiels als Orientierung. Prüfen Sie Ihre zugewiesene Kurve selbst, nennen Sie die tatsächlich verwendeten Quellen und formulieren Sie eigene Schlussfolgerungen. Abweichende Ergebnisse sind zulässig, wenn sie nachvollziehbar belegt werden.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
