from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
from pathlib import Path
from urllib.parse import quote, urlencode

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
EXPORT = ROOT / "data" / "gwl" / "blc-curve-export-v1.json"
SOURCE_TEMPLATE = ROOT / "docs" / "kowo-26-27" / "Arbeitsvorlage_BLC26_Kurvenpruefung.docx"

WORKSET_SCRIPT = Path(__file__).with_name("build-kowo-student-workset.py")
SPEC = importlib.util.spec_from_file_location("kowo_workset", WORKSET_SCRIPT)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Generator kann nicht geladen werden: {WORKSET_SCRIPT}")
WORKSET = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(WORKSET)
body = WORKSET.body
heading = WORKSET.heading
page_title = WORKSET.page_title
run_style = WORKSET.run_style
setup_doc = WORKSET.setup_doc
table = WORKSET.table

SELECTED_ITEMS = [
    "global-temperature",
    "global-warming",
    "radiative-forcing",
    "global-sea-level-rise",
    "arctic-september-sea-ice-area",
    "blue-water-streamflow",
    "green-water-rootzone-soil-moisture",
    "global-forest-status",
    "plastics-microplastics",
    "pfas-pope-global-emissions",
]


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def gwl_url(curve):
    return "https://blcdetlef.github.io/gwl-panel/?" + urlencode(
        {"boundary": curve["boundaryId"], "item": curve["itemId"]}
    )


def blc_url(curve):
    return "https://blcdetlef.github.io/BLC26/?curve=" + quote(curve["curveId"], safe="")


def create_overview(path: Path):
    doc = setup_doc("KoWo 2026/27 · Startanleitung · Version 0.1", compact=True)
    page_title(
        doc,
        "Bitte zuerst lesen",
        "Diese Anleitung zeigt Ihnen, wo Sie beginnen, welche Datei Sie bearbeiten und wo Sie Ihr Ergebnis speichern.",
    )
    heading(doc, "Ihr Arbeitsordner", 1)
    body(doc, "Arbeiten Sie nur im Kurvenordner Ihrer Gruppe unter 01 Kurvenarbeiten. Dort finden Sie START HIER und eine eigene Arbeitsvorlage. Die unveränderten gemeinsamen Unterlagen liegen unter 00 Arbeitsmaterial.")
    heading(doc, "Die ersten Schritte", 1)
    table(doc, ["Schritt", "Was Sie tun"], [
        ["1", "Öffnen Sie in Ihrem Kurvenordner START_HIER.docx und danach beide dort angegebenen Links."],
        ["2", "Lesen Sie die Aufgabenbeschreibung und halten Sie die Schnellhilfe beim Arbeiten geöffnet."],
        ["3", "Tragen Sie eine Gruppenkennung ein. Persönliche Namen sind freiwillig."],
        ["4", "Benennen Sie Arbeitsvorlage.docx in Pruefbericht_<Gruppenkennung>.docx um."],
        ["5", "Bestimmen Sie eine Person, die diese Datei führt. Vermeiden Sie mehrere voneinander abweichende Fassungen."],
        ["6", "Speichern Sie verwendete Studien oder Datenauszüge im Unterordner Quellen und tragen Sie die Fundstellen in den Prüfbericht ein."],
    ], widths=[2.0, 15.1], font_size=8.3)
    heading(doc, "Wenn etwas unklar bleibt", 1)
    body(doc, "Raten Sie nicht. Kennzeichnen Sie den Punkt als nicht prüfbar oder nicht anwendbar, dokumentieren Sie Ihren bisherigen Prüfweg und formulieren Sie im Feld D 02 eine konkrete Frage an die Lehrperson.")
    heading(doc, "Abgabe am Donnerstag", 1)
    body(doc, "Im Kurvenordner liegen der ausgefüllte Prüfbericht und die verwendeten Quellen. Seite 1 enthält den vollständigen studentischen ZUSTAND-Prüfentwurf. Die fachliche Freigabe und das spätere Zertifikat übernimmt die Projektleitung.")
    heading(doc, "Wichtig", 1)
    body(doc, "Bearbeiten oder löschen Sie keine Dateien in anderen Kurvenordnern. Die Ordner 02 Gemeinsame Praesentation und 03 Freigabe und Zertifikate werden nur nach Aufforderung verwendet.")
    doc.save(path)


def create_start_sheet(path: Path, curve):
    doc = setup_doc("KoWo 2026/27 · Kurvenstartblatt · Version 0.1", compact=True)
    page_title(doc, "Startblatt für Ihre Kurvenprüfung", "Diese Kurve ist Ihrer Gruppe als Ausgangspunkt zugeordnet. Prüfen Sie die Inhalte eigenständig und dokumentieren Sie Abweichungen offen.")
    table(doc, ["Merkmal", "Angabe"], [
        ["Kurve im GWL-Menü", curve["label"]],
        ["Bereich", curve["domainLabel"]],
        ["Gruppenkennung", "Bitte eintragen"],
        ["Kurvenkennung", curve["curveId"]],
    ], widths=[4.3, 12.8], font_size=8.0)
    heading(doc, "Geprüfte Startlinks", 1)
    p = doc.add_paragraph()
    run_style(p.add_run("GWL-Beitrag: "), bold=True)
    add_hyperlink(p, "Beitrag im GWL-Panel öffnen", gwl_url(curve))
    p = doc.add_paragraph()
    run_style(p.add_run("BLC26-Kurve: "), bold=True)
    add_hyperlink(p, "Kurve direkt in BLC26 öffnen", blc_url(curve))
    heading(doc, "So arbeiten Sie weiter", 1)
    for text in [
        "Lesen Sie zuerst den GWL-Beitrag und öffnen Sie dort die genannten Quellen.",
        "Prüfen Sie anschließend die Darstellung in BLC26 und bearbeiten Sie Arbeitsvorlage.docx.",
        "Speichern Sie heruntergeladene Quellen im Unterordner Quellen und notieren Sie genaue Fundstellen.",
        "Füllen Sie die ZUSTAND-Seite erst aus, nachdem die fachliche Prüfung abgeschlossen ist.",
    ]:
        p = doc.add_paragraph(style="List Number")
        run_style(p.add_run(text))
    heading(doc, "Falls ein Link oder eine Quelle nicht funktioniert", 1)
    body(doc, "Notieren Sie Link, Zeitpunkt und Ergebnis Ihres Zugriffsversuchs. Bewerten Sie den Punkt grau, statt eine fehlende Information zu erraten, und sammeln Sie die Frage im Prüfbericht unter D 02.")
    doc.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--target-root", required=True, type=Path)
    args = parser.parse_args()
    target_root = args.target_root.resolve()
    curves_root = target_root / "01_Kurvenarbeiten"
    if not target_root.is_dir() or not curves_root.is_dir():
        raise SystemExit(f"Vorbereitete Ordnerstruktur fehlt: {target_root}")

    payload = json.loads(EXPORT.read_text(encoding="utf-8"))
    by_item = {}
    for curve in payload["curves"]:
        by_item.setdefault(curve.get("itemId"), []).append(curve)

    selected = []
    for item_id in SELECTED_ITEMS:
        matches = by_item.get(item_id, [])
        if len(matches) != 1:
            raise SystemExit(f"Keine eindeutige freigegebene Kurve für {item_id}")
        curve = matches[0]
        for key in ("label", "curveId", "boundaryId", "itemId", "domainLabel"):
            if not curve.get(key):
                raise SystemExit(f"Fehlendes Feld {key} bei {item_id}")
        selected.append(curve)

    if len(selected) != 10:
        raise SystemExit("Es müssen genau zehn Kurven ausgewählt sein.")

    create_overview(target_root / "BITTE_ZUERST_LESEN.docx")
    for curve in selected:
        folder = curves_root / curve["label"]
        if not folder.is_dir():
            raise SystemExit(f"Kurvenordner fehlt: {folder}")
        create_start_sheet(folder / "START_HIER.docx", curve)
        shutil.copy2(SOURCE_TEMPLATE, folder / "Arbeitsvorlage.docx")
        (folder / "Quellen").mkdir(exist_ok=True)
        print(curve["itemId"])


if __name__ == "__main__":
    main()
